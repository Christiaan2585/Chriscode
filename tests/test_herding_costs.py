"""Every client's own herding program with costs, like the business's
"Ent en doseer kostes" sheet: five fixed animal groups, lines for one or
more groups, cost of what's used and of whole packs, a medicine box, step
subtotals, a grand total and the cost per animal per year. The sheet below
is made up - the real one is the business's and stays out of this repo."""
import io
import unittest
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlmodel import Session, SQLModel, create_engine, select

from app.api import business, programs
from app.core import herding
from app.models.animal import Animal  # noqa: F401 - registers referenced tables
from app.models.business import BusinessSettings
from app.models.client import Client
from app.models.herd import Herd  # noqa: F401
from app.models.product import Product
from app.models.program import AnimalGroup, HerdingProgram, ProgramStep, ProgramStepProduct, ProgramStepProgress
from app.models.quote_item import QuoteItem
from app.models.user import User


def cost_sheet():
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "Rep Name", "VEEARTSENYPROGRAM KOSTE BEREKENING VIR SKAPE EN BOKKE 2026"
    ws["A4"], ws["B4"] = "=SUM(A5:A9)", "TOTALE DIERE"
    for row, (count, label) in enumerate([(100, "TOTAAL OOIE"), (4, "TOTAAL RAMME"), (110, "TOTAAL LAMMERS "),
                                           (10, "TOTAAL JONG OOITJIES"), (5, "TOTAAL JONG RAMMETJIES")], start=5):
        ws[f"A{row}"], ws[f"B{row}"] = count, label
    for col, head in zip("ABCDEFGH", ["DATUM", "TYD", "PRODUK ", "VERPAK", " PRYS  EXCL VAT ", "DOSERING ml / Dier",
                                      "PRODUK TOTAAL", "TOTAAL R EXCL VAT"]):
        ws[f"{col}10"] = head
    ws["A11"] = "VERANDER NET HIERDIE DATUM"
    ws["A12"], ws["B12"] = datetime(2025, 12, 1), "DEKTYD / LAPA"
    ws["A14"], ws["B14"] = "=A12+147", "LAMTYD"

    def line(row, category, product, dose, groups):
        ws[f"B{row}"], ws[f"C{row}"], ws[f"F{row}"] = category, product, dose
        ws[f"G{row}"] = f'=IFERROR(({groups})*F{row}/D{row},"")' if groups else f'=IFERROR(F{row},"")'

    ws["A16"], ws["B16"] = "=A14+28", "1STE ENTING (1ste Maand)"
    ws["A17"] = "SIT STERTE AF"
    line(17, "ENTING", "Vaccine A - 250ml", 1.5, "$A$7")
    line(19, "MINERALE", "Mineral B 25 Lt", 10, "$A$7")
    ws["B23"] = "TOTAAL"
    ws["A25"], ws["B25"] = "=A14+56", "2de ENTING (2de Maand)"
    line(26, "ENTING", "Vaccine A - 250ml", 2, "$A$7")
    ws["B40"] = "JONG OOITJIES EN RAMME"
    ws["A41"] = "=A74"
    ws["A42"], ws["B42"] = datetime(2026, 3, 24), "Behandel jong ooitjies saam met ooie voor lam"
    line(43, "ENTING", "Vaccine A - 250ml", 2, "$A$8+$A$9")
    ws["A50"] = "=A25+84"
    line(50, "PASTEURELLA + CLOSTRIDIUM", "Vaccine A - 250ml", 2, "$A$8+$A$9")
    ws["A54"] = "=A50"
    line(54, "CHLAMYDIA - OOITJIES", "Vaccine Z (not in the app)", 2, "$A$8")
    ws["A55"] = "=A54"
    line(55, "REV1 - RAMME", "Vaccine A - 250ml", 1, "$A$9")
    ws["A66"], ws["B66"] = "=A12-42", "DOSERING OOIE EN RAMME"
    line(67, "(VOOR DEK ENTING)", "Vaccine A - 250ml", 2, "$A$5+$A$6")
    ws["A74"], ws["B74"] = "=A14-42", "DOSERING OOIE EN RAMME"
    ws["A75"] = datetime(2026, 3, 24)
    line(75, "(VOOR LAM DOSERING)", "Drench C - 5 Lt", 10, "$A$5+$A$6")
    ws["B91"] = "Medisyne Boks:"
    line(92, "", "Mineral B 25 Lt", 2, None)
    ws["B104"] = "TOTALE KOSTE:"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class CostTestCase(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        SQLModel.metadata.create_all(self.engine)
        self.s = Session(self.engine)
        settings = {**BusinessSettings().model_dump(exclude={"id"}), "trading_name": "Sandveld", "vat_registered": True}
        business.update_business(business.BusinessSettingsUpdate(**settings), session=self.s)
        self.client = Client(name="Oom Piet")
        self.user = User(name="Rep", email="rep@example.com", password_hash="x", phone="082 000 0000")
        self.vaccine = Product(name="Vaccine A - 250ml", price=455.69, price_excl_vat=396.25, unit="ml", pack_size=250)
        self.mineral = Product(name="Mineral B 25 Lt", price=7801.31, price_excl_vat=6783.75, unit="ml", pack_size=25000)
        self.drench = Product(name="Drench C - 5 Lt", price=5315.88, price_excl_vat=4622.5, unit="ml", pack_size=5000)
        self.s.add_all([self.client, self.user, self.vaccine, self.mineral, self.drench])
        self.s.commit()
        self.template = herding.get_template(self.s)

    def tearDown(self):
        self.s.close()
        self.engine.dispose()

    def program(self, counts=None):
        program = HerdingProgram(name="Kudde 2026", client_id=self.client.id, mating_date=datetime(2025, 12, 1))
        self.s.add(program)
        self.s.commit()
        for label, count in (counts or {"Ooie": 1200, "Ramme": 48, "Lammers": 1320,
                                        "Jong ooitjies": 140, "Jong rammetjies": 40}).items():
            self.s.add(AnimalGroup(program_id=program.id, animal_type=label, group_size=count))
        self.s.commit()
        return program


class MoneyTests(CostTestCase):
    def test_several_groups_on_one_line(self):
        groups = [AnimalGroup(program_id=1, animal_type=l, group_size=n)
                  for l, n in (("Ooie", 1200), ("Ramme", 48), ("Jong ooitjies", 140), ("Jong rammetjies", 40))]
        self.assertEqual(herding.headcount(groups, "Ooie, Ramme"), 1248)
        self.assertEqual(herding.headcount(groups, "Jong ooitjies + Jong rammetjies"), 180)
        self.assertEqual(herding.headcount(groups, "jong ooitjies"), 140)

    def test_cost_of_what_is_used_and_of_whole_packs(self):
        costs = herding.line_costs(self.vaccine, 1.5, 1320, None, 396.25)
        self.assertEqual((costs["total"], costs["used"], costs["buy"]), (1980.0, 7.92, 8))
        self.assertEqual((costs["cost_used"], costs["cost_buy"]), (3138.30, 3170.00))

    def test_medicine_box_is_a_fixed_quantity(self):
        costs = herding.line_costs(self.mineral, None, 999, 2, 562.5)
        self.assertEqual((costs["used"], costs["buy"], costs["cost_used"]), (2, 2, 1125.0))

    def test_price_excl_vat(self):
        self.assertEqual(herding.price_excl_vat(self.vaccine, business.get_business(self.s)), 396.25)
        no_excl = Product(name="x", price=115.0)
        self.assertEqual(herding.price_excl_vat(no_excl, business.get_business(self.s)), 100.0)


class CostSheetImportTests(CostTestCase):
    def test_sheet_becomes_master_steps_with_lines(self):
        result = herding.import_cost_sheet(self.s, cost_sheet())
        self.assertEqual(result["unmatched"], ["Vaccine Z (not in the app)"])
        # Like the sheet: each block is its own step in its section, even
        # when two sections treat on the same day (lambing -42 here).
        steps = {(s.section, s.anchor, s.offset_days): s for s in herding.ordered_steps(self.s)}
        LAM, JONG, OOI, BOX = herding.SECTIONS
        self.assertEqual(set(steps), {(LAM, "lambing_start", 28), (LAM, "lambing_start", 56),
                                      (JONG, "lambing_start", -42), (JONG, "lambing_start", 140),
                                      (OOI, "mating_start", -42), (OOI, "lambing_start", -42), (BOX, "none", 0)})
        lines = lambda key: [(l.animal_group, l.dose, l.fixed_quantity, l.category) for l in self.s.exec(
            select(ProgramStepProduct).where(ProgramStepProduct.step_id == steps[key].id).order_by(ProgramStepProduct.id)).all()]
        self.assertEqual(lines((LAM, "lambing_start", 28)), [("Lammers", 1.5, None, "Enting"), ("Lammers", 10, None, "Minerale")])
        self.assertEqual(lines((JONG, "lambing_start", 140)), [("Jong ooitjies, Jong rammetjies", 2, None, "Pasteurella + clostridium"),
                                                               ("Jong rammetjies", 1, None, "Rev1 - ramme")])
        self.assertEqual(lines((JONG, "lambing_start", -42)), [("Jong ooitjies, Jong rammetjies", 2, None, "Enting")])
        self.assertEqual(lines((OOI, "lambing_start", -42)), [("Ooie, Ramme", 10, None, "Voor lam dosering")])
        self.assertEqual(lines((BOX, "none", 0)), [(None, None, 2, None)])
        self.assertEqual(steps[(BOX, "none", 0)].stage, "Medisyne boks")
        self.assertIn("Sit sterte af", steps[(LAM, "lambing_start", 28)].management)

    def test_kudde_program_step_on_the_same_day_stays_apart(self):
        kudde = ProgramStep(sort_order=1, anchor="mating_start", offset_days=-42, stage="6 weke voor paring", origin="kudde")
        self.s.add(kudde)
        self.s.commit()
        herding.import_cost_sheet(self.s, cost_sheet())
        same_day = {s.id: s for s in herding.ordered_steps(self.s) if (s.anchor, s.offset_days) == ("mating_start", -42)}
        self.assertEqual(same_day[kudde.id].stage, "6 weke voor paring")  # untouched, still there
        self.assertEqual(sorted(s.section or "" for s in same_day.values()), ["", "Ooie en ramme"])

    def test_reimport_does_not_duplicate(self):
        herding.import_cost_sheet(self.s, cost_sheet())
        before = len(self.s.exec(select(ProgramStepProduct)).all())
        herding.import_cost_sheet(self.s, cost_sheet())
        self.assertEqual(len(self.s.exec(select(ProgramStepProduct)).all()), before)
        self.assertEqual(len(herding.ordered_steps(self.s)), 7)

    def test_dated_heading_without_products_is_still_a_step(self):
        wb = load_workbook(io.BytesIO(cost_sheet()))
        wb.active["A60"], wb.active["B60"] = "=A50+56", "1STE BLOU TONG ENTING"
        buf = io.BytesIO()
        wb.save(buf)
        herding.import_cost_sheet(self.s, buf.getvalue())
        step = next(s for s in herding.ordered_steps(self.s) if (s.anchor, s.offset_days) == ("lambing_start", 196))
        self.assertEqual(step.stage, "1ste blou tong enting")

    def test_not_a_cost_sheet(self):
        with self.assertRaises(herding.TemplateImportError):
            herding.import_cost_sheet(self.s, b"nope")


class ClientProgramTests(CostTestCase):
    def setUp(self):
        super().setUp()
        herding.import_cost_sheet(self.s, cost_sheet())
        self.prog = self.program()

    def schedule(self, today=date(2025, 11, 1)):
        return herding.schedule(self.s, self.prog, today=today)

    def own_lines(self):
        return self.s.exec(select(ProgramStepProduct).join(ProgramStep, ProgramStep.id == ProgramStepProduct.step_id)
                           .where(ProgramStep.program_id == self.prog.id)).all()

    def test_client_gets_their_own_copy(self):
        out = self.schedule()
        self.assertEqual(len(out["steps"]), 7)
        self.assertTrue(all(s["program_id"] == self.prog.id for s in out["steps"]))
        self.assertEqual(len(herding.ordered_steps(self.s)), 7)  # master untouched

    def test_changing_the_client_copy_leaves_the_master_alone(self):
        self.schedule()
        line = self.own_lines()[0]
        programs.update_program_line(self.prog.id, line.id, programs.StepProductFields(dose=9), session=self.s)
        master_doses = {l.dose for l in self.s.exec(select(ProgramStepProduct).join(ProgramStep, ProgramStep.id == ProgramStepProduct.step_id)
                                                     .where(ProgramStep.program_id.is_(None))).all()}
        self.assertNotIn(9, master_doses)

    def test_step_and_grand_totals_and_cost_per_animal(self):
        out = self.schedule()
        first = next(s for s in out["steps"] if (s["anchor"], s["offset_days"]) == ("lambing_start", 28))
        vaccine = first["products"][0]
        self.assertEqual((vaccine["used"], vaccine["buy"], vaccine["cost_used"]), (7.92, 8, 3138.30))
        self.assertAlmostEqual(first["subtotal"], round(3138.30 + 1320 * 10 / 25000 * 6783.75, 2), places=2)
        totals = out["totals"]
        self.assertEqual(totals["animals"], 2748)
        self.assertAlmostEqual(totals["cost_used"], round(sum(s["subtotal"] for s in out["steps"]), 2), places=2)
        self.assertAlmostEqual(totals["per_animal"], round(totals["cost_used"] / 2748, 2), places=2)

    def test_medicine_box_has_no_date_and_is_not_a_step_to_tick(self):
        out = self.schedule()
        box = next(s for s in out["steps"] if s["anchor"] == "none")
        self.assertIsNone(box["date"])
        self.assertEqual(box["status"], "none")
        self.assertEqual(out["steps"][-1]["anchor"], "none")
        self.assertEqual(out["progress"]["total"], 6)

    def test_typed_in_date_wins(self):
        self.schedule()
        step = next(s for s in herding.ordered_steps(self.s, self.prog.id) if s.offset_days == -42 and s.anchor == "lambing_start")
        programs.update_program_step(self.prog.id, step.id, programs.StepFields(date_override=date(2026, 3, 24)), session=self.s)
        moved = next(s for s in self.schedule()["steps"] if s["id"] == step.id)
        self.assertEqual(moved["date"], date(2026, 3, 24))

    def test_old_ticks_carry_over_into_the_copy(self):
        master = next(s for s in herding.ordered_steps(self.s) if s.anchor == "mating_start")
        legacy = self.program()
        self.s.add(ProgramStepProgress(program_id=legacy.id, step_id=master.id))
        self.s.commit()
        out = herding.schedule(self.s, legacy, today=date(2025, 11, 1))
        self.assertEqual(next(s for s in out["steps"] if s["anchor"] == "mating_start")["status"], "done")

    def test_start_again_from_the_master_keeps_ticks(self):
        self.schedule()
        step = next(s for s in herding.ordered_steps(self.s, self.prog.id) if s.anchor == "mating_start")
        programs.set_step_done(self.prog.id, step.id, programs.StepDone(done=True), session=self.s)
        programs.update_program_line(self.prog.id, self.own_lines()[0].id, programs.StepProductFields(dose=9), session=self.s)
        programs.reset_from_master(self.prog.id, session=self.s)
        out = self.schedule()
        self.assertEqual(next(s for s in out["steps"] if s["anchor"] == "mating_start")["status"], "done")
        self.assertNotIn(9, [l.dose for l in self.own_lines()])

    def test_quote_uses_whole_packs_and_the_medicine_box_quantity(self):
        self.schedule()
        lines = self.own_lines()
        box = next(l for l in lines if l.fixed_quantity)
        first_vaccine = next(l for l in lines if l.dose == 1.5)
        result = programs.quote_from_program(self.prog.id, programs.ProgramQuoteRequest(line_ids=[box.id, first_vaccine.id]),
                                             session=self.s, user=self.user)
        items = self.s.exec(select(QuoteItem).where(QuoteItem.quote_id == result["quote"].id)).all()
        self.assertEqual(sorted((i.product_id, i.quantity) for i in items),
                         sorted([(self.mineral.id, 2), (self.vaccine.id, 8)]))

    def test_cost_pdf(self):
        response = programs.program_cost_pdf(self.prog.id, session=self.s)
        self.assertTrue(response.body.startswith(b"%PDF"))

    def test_deleting_the_program_removes_its_copy(self):
        self.schedule()
        programs.delete_program(self.prog.id, session=self.s)
        self.assertEqual(herding.ordered_steps(self.s, self.prog.id), [])
        self.assertEqual(self.own_lines(), [])

    def test_head_counts(self):
        programs.set_head_counts(self.prog.id, programs.HeadCounts(counts={"ooie": 1300, "Jong rammetjies": 0,
                                                                            "Bokke": 12}), session=self.s)
        counts = {g.animal_type: g.group_size for g in self.s.exec(
            select(AnimalGroup).where(AnimalGroup.program_id == self.prog.id)).all()}
        self.assertEqual(counts, {"Ooie": 1300, "Ramme": 48, "Lammers": 1320, "Jong ooitjies": 140, "Bokke": 12})

    def test_master_preview_dates_skip_the_medicine_box(self):
        out = programs.read_template(mating_date=date(2025, 12, 1), session=self.s)
        box = next(s for s in out["steps"] if s["anchor"] == "none")
        self.assertIsNone(box["date"])

    def test_line_for_another_program_is_refused(self):
        self.schedule()
        other = self.program()
        herding.schedule(self.s, other, today=date(2025, 11, 1))
        line = self.own_lines()[0]
        with self.assertRaises(HTTPException) as ctx:
            programs.update_program_line(other.id, line.id, programs.StepProductFields(dose=1), session=self.s)
        self.assertEqual(ctx.exception.status_code, 404)


if __name__ == "__main__":
    unittest.main()
