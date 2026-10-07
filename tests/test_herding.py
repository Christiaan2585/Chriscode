import io
import unittest
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import Workbook
from sqlmodel import Session, SQLModel, create_engine, select

from app.api import programs
from app.core import cascade, herding
from app.models.animal import Animal  # noqa: F401 - registers referenced tables
from app.models.client import Client
from app.models.herd import Herd  # noqa: F401
from app.models.product import Product
from app.models.program import (AnimalGroup, HerdingProgram, ProgramStep, ProgramStepProduct,
                                ProgramStepProgress, ProgramTemplate)
from app.models.quote_item import QuoteItem
from app.models.user import User


def sheet_like_the_business_uses():
    """Same layout and date formulas as the business's own sheet; the text is made up."""
    wb = Workbook()
    ws = wb.active
    ws.title = "PROGRAM"
    ws["A1"], ws["B1"] = "JAAR", 2025
    ws["E1"], ws["G1"] = "EERSTE PAARDATUM =", datetime(2025, 12, 1)
    ws["A2"], ws["G2"] = "LENGTE VAN PAARSEISOEN (WEKE) =", 6
    ws["A3"], ws["G3"] = "OUDERDOM WAAROP LAMMERS GESPEEN WORD (MAANDE) =", 4
    for col, head in zip("ABCDEFG", ["Datum", "Produksie Stadium", "Kudde-Bestuur", "Inentings",
                                     "Doserings", "Vitamiene en Spoor-elemente", "Voeding"]):
        ws[f"{col}4"] = head
    ws["A5"] = "EERSTE LAMSEISOEN"
    rows = {
        6: ("=$G1-(8*7)", "8 weeks before", "1  Check rams                2  Trim hooves",
            "Vaccine X  3ml under the skin          Vaccine W 3 (if needed)", None),
        11: ("=G1-(7*6)", "6 weeks before", "1  Crutch ewes", None, None),
        21: ("=$G$1", "MATING STARTS", "1  Rams in", None, None),
        26: ("=$G$1+(G2*7)", "MATING ENDS", "1  Rams out", None, None),
        31: ("=A26+28", None, None, "Booster B", None),
        46: ("=+A51-(7*6)", "6 weeks before lambing", None, None, "1  Ewes - worms"),
        51: ("=G1+147", "Lambing starts", "1  Ewes lamb", None, None),
        61: ("=+A51+(($G$2*7)+3)", None, "1  Lambing done", None, None),
        76: ("=A21+150+90", "Weaning day", "1  Speen lammers", None, None),
        81: ("=+A76+(7*4)", "Replacement ewes", "1. Vaccinate", "1. Vaccine Y", None),
    }
    for row, (formula, stage, management, vaccinations, dosing) in rows.items():
        ws[f"A{row}"], ws[f"B{row}"], ws[f"C{row}"] = formula, stage, management
        ws[f"D{row}"], ws[f"E{row}"] = vaccinations, dosing
    ws["C82"], ws["D82"] = "2. Second vaccine", "2. Vaccine Z"  # a step spread over two rows
    ws["I81"] = "=250*1.4*2"  # scratch sums beside the table are ignored
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class HerdingTestCase(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        SQLModel.metadata.create_all(self.engine)
        self.s = Session(self.engine)
        self.client = Client(name="Oom Piet")
        self.user = User(name="Rep", email="rep@example.com", password_hash="x")
        self.s.add_all([self.client, self.user])
        self.s.commit()
        self.template = herding.get_template(self.s)

    def tearDown(self):
        self.s.close()
        self.engine.dispose()

    def program(self, **fields):
        program = HerdingProgram(name="Kudde 2026", client_id=self.client.id,
                                 mating_date=datetime(2025, 12, 1), **fields)
        self.s.add(program)
        self.s.commit()
        return program


class AnchorDateTests(HerdingTestCase):
    def test_dates_follow_the_first_mating_day(self):
        anchors = herding.anchor_dates(self.program(), self.template)
        self.assertEqual(anchors, {
            "mating_start": date(2025, 12, 1),
            "mating_end": date(2026, 1, 12),       # 6 weeks
            "lambing_start": date(2026, 4, 27),    # + 147 days
            "lambing_end": date(2026, 6, 8),       # lambing + the 6 week season
            "weaning": date(2026, 7, 29),          # the sheet's own rule: mating + 240 days
        })

    def test_weaning_by_age(self):
        anchors = herding.anchor_dates(self.program(weaning_rule="age", weaning_months=4), self.template)
        self.assertEqual(anchors["weaning"], date(2026, 8, 27))  # lambing + 4 months

    def test_program_overrides_the_template(self):
        anchors = herding.anchor_dates(self.program(mating_weeks=5, weaning_days=250), self.template)
        self.assertEqual(anchors["mating_end"], date(2026, 1, 5))
        self.assertEqual(anchors["weaning"], date(2026, 8, 8))

    def test_month_end_is_clamped(self):
        self.assertEqual(herding.add_months(date(2026, 1, 31), 1), date(2026, 2, 28))


class DoseTests(unittest.TestCase):
    def packs(self, product, dose, head):
        costs = herding.line_costs(product, dose, head, None, product.price)
        return costs["total"], costs["buy"]

    def test_bottles_are_rounded_up(self):
        product = Product(name="Vaccine X 100ml", price=100, unit="ml", pack_size=100)
        self.assertEqual(self.packs(product, 2, 120), (240, 3))

    def test_product_without_pack_size_is_counted_per_dose(self):
        product = Product(name="Ear tag", price=10)
        self.assertEqual(self.packs(product, 1, 120), (120, 120))

    def test_no_dose_or_no_animals_means_nothing(self):
        product = Product(name="Vaccine X", price=100, pack_size=100)
        self.assertEqual(self.packs(product, None, 120), (0, 0))
        self.assertEqual(self.packs(product, 2, 0), (0, 0))

    def test_headcount_by_group(self):
        groups = [AnimalGroup(program_id=1, animal_type="Ooie", group_size=120),
                  AnimalGroup(program_id=1, animal_type="Ramme", group_size=4)]
        self.assertEqual(herding.headcount(groups, "ooie"), 120)
        self.assertEqual(herding.headcount(groups, None), 124)
        self.assertEqual(herding.headcount(groups, "Lammers"), 0)


class ScheduleAndQuoteTests(HerdingTestCase):
    def setUp(self):
        super().setUp()
        self.before = ProgramStep(sort_order=1, anchor="mating_start", offset_days=-56, stage="8 weeks before")
        self.lambing = ProgramStep(sort_order=2, anchor="lambing_start", offset_days=0, stage="Lambing")
        self.vaccine = Product(name="Vaccine X", price=250.0, unit="ml", pack_size=100)
        self.drench = Product(name="Drench", price=80.0, unit="ml", pack_size=500)
        self.s.add_all([self.before, self.lambing, self.vaccine, self.drench])
        self.s.commit()
        self.line_ewes = ProgramStepProduct(step_id=self.before.id, product_id=self.vaccine.id, animal_group="Ooie", dose=3)
        self.line_rams = ProgramStepProduct(step_id=self.before.id, product_id=self.vaccine.id, animal_group="Ramme", dose=3)
        self.line_lambs = ProgramStepProduct(step_id=self.lambing.id, product_id=self.drench.id, animal_group="Lammers", dose=5)
        self.s.add_all([self.line_ewes, self.line_rams, self.line_lambs])
        self.prog = self.program()
        self.s.add_all([AnimalGroup(program_id=self.prog.id, animal_type="Ooie", group_size=100),
                        AnimalGroup(program_id=self.prog.id, animal_type="Ramme", group_size=4)])
        self.s.commit()

    def test_schedule_dates_statuses_and_amounts(self):
        out = herding.schedule(self.s, self.prog, today=date(2025, 10, 8))
        first, second = out["steps"]
        self.assertEqual((first["date"], first["status"]), (date(2025, 10, 6), "overdue"))
        self.assertEqual((second["date"], second["status"]), (date(2026, 4, 27), "upcoming"))
        ewes = first["products"][0]
        self.assertEqual((ewes["head"], ewes["total"], ewes["buy"]), (100, 300, 3))
        self.assertEqual(out["progress"], {"done": 0, "total": 2})

    def test_step_due_within_a_week(self):
        out = herding.schedule(self.s, self.prog, today=date(2025, 10, 1))
        self.assertEqual(out["steps"][0]["status"], "due")

    def test_ticking_a_step_off(self):
        programs.set_step_done(self.prog.id, self.before.id, programs.StepDone(done=True), session=self.s)
        out = herding.schedule(self.s, self.prog, today=date(2025, 10, 8))
        self.assertEqual(out["steps"][0]["status"], "done")
        self.assertEqual(out["progress"]["done"], 1)
        programs.set_step_done(self.prog.id, self.before.id, programs.StepDone(done=False), session=self.s)
        self.assertEqual(herding.schedule(self.s, self.prog)["progress"]["done"], 0)

    def test_the_programs_quote_has_whole_packs_per_line(self):
        out = programs.program_schedule(self.prog.id, session=self.s)
        items = self.s.exec(select(QuoteItem).where(QuoteItem.quote_id == out["quote"]["id"])).all()
        # 100 ewes x 3 ml = 3 bottles, 4 rams x 3 ml = 1 bottle; no lambs counted -> no drench on the quote
        self.assertEqual(sorted((i.product_id, i.quantity) for i in items), [(self.vaccine.id, 1), (self.vaccine.id, 3)])
        self.assertEqual(out["quote"]["client_id"], self.client.id)

    def test_program_without_mating_date_has_no_schedule(self):
        legacy = HerdingProgram(name="Old", client_id=self.client.id)
        self.s.add(legacy)
        self.s.commit()
        self.assertEqual(herding.schedule(self.s, legacy)["steps"], [])

    def test_calendar_lists_steps_in_range(self):
        events = programs.program_calendar(start=date(2025, 10, 1), end=date(2025, 10, 31), session=self.s)
        self.assertEqual([(e["date"], e["stage"], e["client_name"]) for e in events],
                         [(date(2025, 10, 6), "8 weeks before", "Oom Piet")])

    def test_farmer_pdf(self):
        response = programs.program_pdf(self.prog.id, session=self.s)
        self.assertTrue(response.body.startswith(b"%PDF"))

    def test_deleting_the_program_removes_its_progress(self):
        programs.set_step_done(self.prog.id, self.before.id, programs.StepDone(done=True), session=self.s)
        programs.delete_program(self.prog.id, session=self.s)
        self.assertEqual(self.s.exec(select(ProgramStepProgress)).all(), [])

    def test_deleting_a_product_unlinks_it_from_the_program(self):
        cascade.delete_product(self.s, self.drench)
        self.s.commit()
        self.assertEqual([l.product_id for l in self.s.exec(select(ProgramStepProduct)).all()],
                         [self.vaccine.id, self.vaccine.id])

    def test_deleting_a_step_removes_its_products_and_progress(self):
        programs.delete_template_step(self.before.id, session=self.s)
        self.assertEqual([l.step_id for l in self.s.exec(select(ProgramStepProduct)).all()], [self.lambing.id])

    def test_deleting_a_master_step_leaves_client_copies_alone(self):
        programs.set_step_done(self.prog.id, self.before.id, programs.StepDone(done=True), session=self.s)
        programs.delete_template_step(self.before.id, session=self.s)
        out = herding.schedule(self.s, self.prog, today=date(2025, 10, 8))
        self.assertEqual([s["status"] for s in out["steps"]], ["done", "upcoming"])
        self.assertEqual(self.s.exec(select(ProgramStepProgress)).all(), [])


class ExcelImportTests(HerdingTestCase):
    def import_sheet(self, data=None):
        return herding.import_template(self.s, data or sheet_like_the_business_uses())

    def steps(self):
        return self.s.exec(select(ProgramStep).order_by(ProgramStep.sort_order)).all()

    def test_formulas_become_date_rules(self):
        self.import_sheet()
        self.assertEqual([(s.anchor, s.offset_days) for s in self.steps()], [
            ("mating_start", -56), ("mating_start", -42), ("mating_start", 0), ("mating_end", 0),
            ("mating_end", 28), ("lambing_start", -42), ("lambing_start", 0), ("lambing_end", 3),
            ("weaning", 0), ("weaning", 28),
        ])

    def test_settings_come_from_the_sheet(self):
        self.import_sheet()
        t = herding.get_template(self.s)
        self.assertEqual((t.mating_weeks, t.weaning_months, t.gestation_days, t.weaning_days, t.weaning_rule),
                         (6, 4, 147, 240, "fixed"))

    def test_imported_dates_match_the_sheet(self):
        self.import_sheet()
        program = self.program()
        dates = [s["date"] for s in herding.schedule(self.s, program)["steps"]]
        self.assertEqual(dates[0], date(2025, 10, 6))
        self.assertEqual(dates[-2], date(2026, 7, 29))  # weaning, as the sheet works it out
        self.assertEqual(dates[-1], date(2026, 8, 26))

    def test_cell_text_is_tidied_and_rows_joined(self):
        self.import_sheet()
        first, last = self.steps()[0], self.steps()[-1]
        self.assertEqual(first.management, "1 Check rams\n2 Trim hooves")
        # a long run of spaces is how the sheet starts a new line in a cell
        self.assertEqual(first.vaccinations, "Vaccine X 3ml under the skin\nVaccine W 3 (if needed)")
        self.assertEqual(last.management, "1. Vaccinate\n2. Second vaccine")
        self.assertEqual(last.vaccinations, "1. Vaccine Y\n2. Vaccine Z")
        self.assertEqual(first.stage, "8 weeks before")

    def test_reimport_keeps_products_and_progress_on_matching_steps(self):
        self.import_sheet()
        step = self.steps()[0]
        product = Product(name="Vaccine X", price=1)
        self.s.add(product)
        self.s.commit()
        self.s.add(ProgramStepProduct(step_id=step.id, product_id=product.id, dose=3))
        self.s.commit()
        result = self.import_sheet()
        self.assertEqual(self.steps()[0].id, step.id)
        self.assertEqual(len(self.s.exec(select(ProgramStepProduct)).all()), 1)
        self.assertEqual((result["created"], result["removed"]), (0, 0))

    def test_not_a_workbook(self):
        with self.assertRaises(herding.TemplateImportError):
            self.import_sheet(b"not excel")

    def test_sheet_without_a_mating_date_is_refused(self):
        wb = Workbook()
        wb.active["A1"] = "Datum"
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(herding.TemplateImportError):
            self.import_sheet(buf.getvalue())


if __name__ == "__main__":
    unittest.main()
