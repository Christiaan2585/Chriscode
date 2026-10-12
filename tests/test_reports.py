"""Client statements, the aged-debtors report and the accountant's Excel export."""
import io
import unittest
from datetime import date, datetime

from openpyxl import load_workbook
from sqlmodel import Session, SQLModel, create_engine

from app import main  # noqa: F401 - registers every table
from app.core import reports
from app.models.business import BusinessSettings
from app.models.client import Client
from app.models.invoice import Invoice, InvoiceItem
from app.models.payment import CreditNote, Payment
from app.models.product import Product


class ReportCase(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        SQLModel.metadata.create_all(self.engine)
        self.s = Session(self.engine)
        self.s.add(BusinessSettings(id=1, trading_name="Test", vat_registered=True, vat_rate=15.0))
        self.jan = Client(name="Jan Smit", farm_name="Rietvlei")
        self.piet = Client(name="Piet", farm_name="=HYPERLINK(\"http://evil\")")
        self.product = Product(name="Dip", price=115.0)
        self.s.add_all([self.jan, self.piet, self.product])
        self.s.commit()

    def tearDown(self):
        self.s.close()
        self.engine.dispose()

    def invoice(self, client, number, when, due, total, status="unpaid"):
        inv = Invoice(client_id=client.id, number=number, date=when, due_date=due, total_amount=total, status=status)
        self.s.add(inv)
        self.s.commit()
        self.s.add(InvoiceItem(invoice_id=inv.id, product_id=self.product.id, quantity=1, unit_price=round(total / 1.15, 2),
                               subtotal=total, vat_percent=15.0))
        self.s.commit()
        return inv


class StatementTests(ReportCase):
    def test_running_balance_with_invoices_payments_and_credits_in_date_order(self):
        a = self.invoice(self.jan, "INV0000001", datetime(2026, 8, 1), datetime(2026, 8, 31), 1150.0)
        self.invoice(self.jan, "INV0000002", datetime(2026, 9, 10), datetime(2026, 10, 10), 230.0)
        self.s.add(Payment(invoice_id=a.id, client_id=self.jan.id, amount=500.0, date=datetime(2026, 8, 20), method="EFT", reference="ref 9"))
        self.s.add(CreditNote(number="CRN0000001", invoice_id=a.id, client_id=self.jan.id, amount=115.0, vat_amount=15.0,
                              reason="Returned", date=datetime(2026, 9, 1)))
        self.s.commit()
        st = reports.client_statement(self.s, self.jan.id, date(2026, 8, 1), date(2026, 9, 30))
        self.assertEqual([r["reference"] for r in st["rows"]], ["INV0000001", "Payment", "CRN0000001", "INV0000002"])
        self.assertEqual([r["balance"] for r in st["rows"]], [1150.0, 650.0, 535.0, 765.0])
        self.assertEqual((st["opening_balance"], st["closing_balance"]), (0.0, 765.0))

    def test_what_came_before_the_period_is_the_opening_balance(self):
        self.invoice(self.jan, "INV0000001", datetime(2026, 7, 1), datetime(2026, 7, 31), 100.0)
        self.invoice(self.jan, "INV0000002", datetime(2026, 9, 5), datetime(2026, 10, 5), 50.0)
        st = reports.client_statement(self.s, self.jan.id, date(2026, 9, 1), date(2026, 9, 30))
        self.assertEqual(st["opening_balance"], 100.0)
        self.assertEqual(st["closing_balance"], 150.0)
        self.assertEqual(len(st["rows"]), 1)

    def test_cancelled_invoices_and_other_clients_are_left_out(self):
        self.invoice(self.jan, "INV0000001", datetime(2026, 9, 1), datetime(2026, 9, 30), 100.0, status="cancelled")
        self.invoice(self.piet, "INV0000002", datetime(2026, 9, 1), datetime(2026, 9, 30), 999.0)
        st = reports.client_statement(self.s, self.jan.id, date(2026, 9, 1), date(2026, 9, 30))
        self.assertEqual(st["rows"], [])
        self.assertEqual(st["closing_balance"], 0.0)

    def test_the_statement_pdf_is_a_real_pdf_naming_the_client_and_the_amount_due(self):
        from pypdf import PdfReader
        self.invoice(self.jan, "INV0000001", datetime(2026, 9, 1), datetime(2026, 9, 30), 1150.0)
        pdf = reports.statement_pdf(self.s, self.jan.id, date(2026, 9, 1), date(2026, 9, 30)).getvalue()
        text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
        self.assertIn("STATEMENT", text)
        self.assertIn("INV0000001", text)
        self.assertIn("1,150.00", text)


class AgedDebtorTests(ReportCase):
    def test_outstanding_amounts_land_in_the_right_age_bucket(self):
        today = date(2026, 10, 12)
        self.invoice(self.jan, "INV1", datetime(2026, 10, 1), datetime(2026, 10, 31), 100.0)   # not due yet
        self.invoice(self.jan, "INV2", datetime(2026, 9, 1), datetime(2026, 10, 1), 200.0)    # 11 days late
        self.invoice(self.jan, "INV3", datetime(2026, 7, 1), datetime(2026, 8, 1), 300.0)     # 72 days late
        self.invoice(self.jan, "INV4", datetime(2026, 5, 1), datetime(2026, 6, 1), 400.0)     # 133 days late
        paid = self.invoice(self.piet, "INV5", datetime(2026, 9, 1), datetime(2026, 9, 2), 50.0)
        self.s.add(Payment(invoice_id=paid.id, client_id=self.piet.id, amount=50.0, date=datetime(2026, 9, 3)))
        self.s.commit()
        out = reports.aged_debtors(self.s, today)
        self.assertEqual(len(out["clients"]), 1)  # Piet paid up: not a debtor
        row = out["clients"][0]
        self.assertEqual((row["current"], row["days_30"], row["days_60"], row["days_90"], row["older"], row["total"]),
                         (100.0, 200.0, 0.0, 300.0, 400.0, 1000.0))
        self.assertEqual(out["totals"]["total"], 1000.0)


class AccountingExportTests(ReportCase):
    def test_invoices_payments_credit_notes_and_vat_in_the_period(self):
        a = self.invoice(self.jan, "INV0000001", datetime(2026, 9, 1), datetime(2026, 9, 30), 1150.0)
        self.invoice(self.jan, "INV0000002", datetime(2026, 11, 1), datetime(2026, 11, 30), 500.0)  # outside the period
        self.s.add(Payment(invoice_id=a.id, client_id=self.jan.id, amount=1000.0, date=datetime(2026, 9, 15), method="EFT"))
        self.s.add(CreditNote(number="CRN0000001", invoice_id=a.id, client_id=self.jan.id, amount=115.0, vat_amount=15.0,
                              reason="Returned", date=datetime(2026, 9, 20)))
        self.s.commit()
        wb = load_workbook(io.BytesIO(reports.accounting_workbook(self.s, date(2026, 9, 1), date(2026, 9, 30))))
        self.assertEqual(wb.sheetnames, ["Invoices", "Payments", "Credit notes", "VAT"])
        invoices = list(wb["Invoices"].values)
        self.assertEqual(len(invoices), 2)  # header + the one invoice in September
        self.assertEqual(invoices[1][0], "INV0000001")
        self.assertEqual(list(wb["Payments"].values)[1][-1], 1000.0)
        vat = {r[0]: r for r in wb["VAT"].values}
        self.assertEqual(vat["2026-09"][1:], (150.0, 15.0, 135.0))  # output VAT, VAT credited, net

    def test_text_that_looks_like_a_formula_stays_plain_text(self):
        self.invoice(self.piet, "INV0000009", datetime(2026, 9, 1), datetime(2026, 9, 30), 100.0)
        wb = load_workbook(io.BytesIO(reports.accounting_workbook(self.s, date(2026, 9, 1), date(2026, 9, 30))))
        cell = [c for c in wb["Invoices"]["D"] if str(c.value).startswith("=")]
        self.assertTrue(cell and all(c.data_type == "s" for c in cell))


class RouteTests(ReportCase):
    def test_the_accountant_export_is_admin_only_and_audited(self):
        from app.api import reports as api
        from app.core.audit import describe
        from app.core.security import require_admin
        route = next(r for r in api.router.routes if r.path.endswith("/accounting.xlsx"))
        self.assertIn(require_admin, [d.dependency for d in route.dependencies])
        self.assertEqual(describe("GET", "/reports/accounting.xlsx")["action"], "Exported the accounts to Excel")

    def test_a_period_that_ends_before_it_starts_is_refused(self):
        from fastapi import HTTPException
        from app.api import reports as api
        with self.assertRaises(HTTPException) as ctx:
            api.statement(self.jan.id, start=date(2026, 9, 30), end=date(2026, 9, 1), session=self.s)
        self.assertEqual(ctx.exception.status_code, 422)

    def test_an_unknown_client_is_404(self):
        from fastapi import HTTPException
        from app.api import reports as api
        with self.assertRaises(HTTPException) as ctx:
            api.statement(9999, session=self.s)
        self.assertEqual(ctx.exception.status_code, 404)


if __name__ == "__main__":
    unittest.main()
