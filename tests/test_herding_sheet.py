"""A client's program as the business's cost sheet: sections with their own
totals, downloaded as an .xlsx in the sheet's layout (with live formulas)
and read back into a client's program. Made-up data, like test_herding_costs."""
import io
import unittest
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlmodel import select

from app.api import programs
from app.core import herding
from app.models.program import AnimalGroup, HerdingProgram, ProgramStep, ProgramStepProduct
from tests.test_herding_costs import CostTestCase, cost_sheet

LAM, JONG, OOI, BOX = herding.SECTIONS


class SheetTestCase(CostTestCase):
    def setUp(self):
        super().setUp()
        herding.import_cost_sheet(self.s, cost_sheet())
        self.prog = self.program()

    def schedule(self, program=None):
        return herding.schedule(self.s, program or self.prog, today=date(2025, 11, 1))

    def export(self, program=None):
        return programs.program_sheet(( program or self.prog).id, session=self.s).body


class SectionTests(SheetTestCase):
    def test_sections_in_sheet_order_with_totals(self):
        out = self.schedule()
        self.assertEqual([s["name"] for s in out["sections"]], [LAM, JONG, OOI, BOX])
        self.assertAlmostEqual(sum(s["subtotal"] for s in out["sections"]), out["totals"]["cost_used"], places=2)
        box = out["sections"][-1]
        self.assertEqual(box["subtotal"], round(2 * 6783.75, 2))

    def test_older_steps_get_a_section_from_their_animals(self):
        step = ProgramStep(program_id=None, anchor="lambing_start", offset_days=10, sort_order=99)
        self.s.add(step)
        self.s.commit()
        self.s.add(ProgramStepProduct(step_id=step.id, product_id=self.vaccine.id, animal_group="Lammers", dose=1))
        self.s.commit()
        fresh = self.program()
        sections = {s["source_step_id"]: s["section"] for s in self.schedule(fresh)["steps"]}
        self.assertEqual(sections[step.id], LAM)


class ExcelOutTests(SheetTestCase):
    def test_download_is_the_sheet_with_live_formulas(self):
        ws = load_workbook(io.BytesIO(self.export())).worksheets[0]
        self.assertEqual(ws["A5"].value, 1200)
        self.assertEqual(ws["A12"].value, datetime(2025, 12, 1))
        self.assertTrue(str(ws["A14"].value).startswith("=A12+"))
        formulas = [c.value for row in ws.iter_rows(min_col=7, max_col=8) for c in row if isinstance(c.value, str)]
        self.assertTrue(any(f.startswith("=IFERROR(") for f in formulas))
        texts = [str(c.value) for row in ws.iter_rows(max_col=4) for c in row if c.value]
        self.assertIn("Vaccine A - 250ml", texts)
        self.assertTrue(any("TOTALE KOSTE" in t for t in texts))


class ExcelInjectionTests(SheetTestCase):
    def test_typed_in_text_is_never_a_formula(self):
        self.client.name = '=HYPERLINK("http://example.com","x")'
        self.s.add(self.client)
        step = next(s for s in herding.ordered_steps(self.s) if s.section == LAM)
        step.stage = "=1+1"
        self.s.add(step)
        self.s.commit()
        ws = load_workbook(io.BytesIO(self.export(self.program()))).worksheets[0]
        cells = {c.value: c.data_type for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
        self.assertEqual(cells[self.client.name], "s")
        self.assertEqual(cells["=1+1"], "s")
        self.assertEqual(cells["=A12+147"], "f")  # the sheet's own formulas still work

    def test_such_text_reads_back_as_text(self):
        step = next(s for s in herding.ordered_steps(self.s) if s.section == LAM)
        step.stage = "=1+1"
        self.s.add(step)
        self.s.commit()
        fresh = self.program()
        data = self.export(fresh)
        programs.import_program_sheet_bytes(fresh.id, data, session=self.s)  # no "couldn't work out the date"


class ExcelInTests(SheetTestCase):
    def other_program(self):
        other = HerdingProgram(name="Nuut", client_id=self.client.id, mating_date=datetime(2026, 1, 5))
        self.s.add(other)
        self.s.commit()
        self.s.add(AnimalGroup(program_id=other.id, animal_type="Ooie", group_size=5))
        self.s.commit()
        return other

    def test_round_trip_gives_the_same_program(self):
        data = self.export()
        other = self.other_program()
        programs.import_program_sheet_bytes(other.id, data, session=self.s)
        self.s.refresh(other)
        a, b = self.schedule(), self.schedule(other)
        self.assertEqual(other.mating_date.date(), date(2025, 12, 1))
        self.assertEqual(b["totals"]["animals"], 2748)
        self.assertAlmostEqual(b["totals"]["cost_used"], a["totals"]["cost_used"], places=2)
        self.assertEqual([s["name"] for s in b["sections"]], [s["name"] for s in a["sections"]])

    def test_a_dose_changed_in_excel_comes_in(self):
        wb = load_workbook(io.BytesIO(self.export()))
        ws = wb.worksheets[0]
        row = next(r for r in range(1, ws.max_row + 1) if ws[f"C{r}"].value == "Drench C - 5 Lt")
        ws[f"F{row}"] = 20
        buf = io.BytesIO()
        wb.save(buf)
        programs.import_program_sheet_bytes(self.prog.id, buf.getvalue(), session=self.s)
        doses = [l["dose"] for s in self.schedule()["steps"] for l in s["products"] if l["product_id"] == self.drench.id]
        self.assertEqual(doses, [20])

    def test_ticks_survive_an_import(self):
        step = next(s for s in self.schedule()["steps"] if s["section"] == LAM)
        programs.set_step_done(self.prog.id, step["id"], programs.StepDone(done=True), session=self.s)
        programs.import_program_sheet_bytes(self.prog.id, self.export(), session=self.s)
        again = next(s for s in self.schedule()["steps"] if (s["section"], s["anchor"], s["offset_days"])
                     == (LAM, step["anchor"], step["offset_days"]))
        self.assertEqual(again["status"], "done")

    def test_not_a_sheet(self):
        with self.assertRaises(HTTPException) as ctx:
            programs.import_program_sheet_bytes(self.prog.id, b"nope", session=self.s)
        self.assertEqual(ctx.exception.status_code, 422)


if __name__ == "__main__":
    unittest.main()
