"""The master herding program: working out each step's date for a client's
program, the product amounts from its headcounts, and importing the
business's Excel sheet (where every date is a formula off the first mating
day) as the template."""
import ast
import calendar
import io
import math
import re
from datetime import date, datetime, timedelta
from types import SimpleNamespace

from openpyxl import load_workbook
from sqlmodel import Session, select

from app.models.business import BusinessSettings
from app.models.product import Product
from app.models.program import (AnimalGroup, HerdingProgram, ProgramStep, ProgramStepProduct,
                                ProgramStepProgress, ProgramTemplate)

ANCHORS = ("mating_start", "mating_end", "lambing_start", "lambing_end", "weaning")
UNDATED = "none"  # a step with no date - the medicine box
TEXT_FIELDS = ("stage", "management", "vaccinations", "dosing", "vitamins", "feeding")
DUE_WITHIN_DAYS = 7
# The business's cost sheet counts animals in these five groups; a product
# line names one or more of them ("Ooie, Ramme").
GROUPS = ("Ooie", "Ramme", "Lammers", "Jong ooitjies", "Jong rammetjies")
# The cost sheet's sections, in its order; each has its own TOTAAL.
SECTIONS = ("Lammers", "Jong ooitjies en ramme", "Ooie en ramme", "Medisyne boks")
OTHER_SECTION = "Ander"  # products for a mix of animals
_SECTION_GROUPS = {SECTIONS[0]: {"lammers"}, SECTIONS[1]: {"jong ooitjies", "jong rammetjies"},
                   SECTIONS[2]: {"ooie", "ramme"}}


def derive_section(lines) -> str | None:
    """A step's section from its lines' animals, for steps made before
    sections existed (or added by hand): the medicine box if it's fixed
    quantities, else the section all its animals belong to."""
    if not lines:
        return None
    if all(l.fixed_quantity for l in lines):
        return SECTIONS[3]
    names = {n for l in lines for n in group_names(l.animal_group)}
    return next((s for s, groups in _SECTION_GROUPS.items() if names and names <= groups), None)


def get_template(session: Session) -> ProgramTemplate:
    template = session.get(ProgramTemplate, 1)
    if template is None:
        template = ProgramTemplate(id=1)
        session.add(template)
        session.commit()
        session.refresh(template)
    return template


def add_months(day: date, months: int) -> date:
    month = day.month - 1 + months
    year, month = day.year + month // 12, month % 12 + 1
    return date(year, month, min(day.day, calendar.monthrange(year, month)[1]))


def anchor_dates(program: HerdingProgram, template: ProgramTemplate) -> dict:
    mating = program.mating_date.date() if isinstance(program.mating_date, datetime) else program.mating_date
    season = timedelta(weeks=program.mating_weeks or template.mating_weeks)
    lambing = mating + timedelta(days=template.gestation_days)
    if (program.weaning_rule or template.weaning_rule) == "age":
        weaning = add_months(lambing, program.weaning_months or template.weaning_months)
    else:
        weaning = mating + timedelta(days=program.weaning_days or template.weaning_days)
    return {"mating_start": mating, "mating_end": mating + season, "lambing_start": lambing,
            "lambing_end": lambing + season, "weaning": weaning}


def group_names(animal_group) -> list:
    """"Ooie, Ramme" or "Jong ooitjies + Jong rammetjies" -> the names."""
    return [n.strip().lower() for n in re.split(r"[,+&]", animal_group or "") if n.strip()]


def headcount(groups, animal_group) -> int:
    wanted = group_names(animal_group)
    return sum(g.group_size for g in groups if not wanted or g.animal_type.strip().lower() in wanted)


def price_excl_vat(product: Product, settings: BusinessSettings) -> float:
    """The product's price before VAT - what the cost sheet works in."""
    if product.price_excl_vat:
        return product.price_excl_vat
    if settings.vat_registered and settings.vat_rate:
        return round(product.price / (1 + settings.vat_rate / 100), 2)
    return product.price


def line_costs(product: Product, dose, head, fixed_quantity, price) -> dict:
    """Like the cost sheet: `used` = packs actually used (7.92 bottles) and
    its cost; `buy` = whole packs (8) and their cost. The medicine box is a
    fixed number of packs whatever the headcount. Without a pack size one
    dose is one unit sold."""
    if fixed_quantity:
        total, exact = None, fixed_quantity
    else:
        total = round((dose or 0) * (head or 0), 3)
        exact = total / product.pack_size if product.pack_size else total
    buy = math.ceil(exact - 1e-9) if exact > 0 else 0
    return {"total": total, "used": round(exact, 2), "buy": buy,
            "cost_used": round(exact * price, 2), "cost_buy": round(buy * price, 2)}


def ordered_steps(session: Session, program_id: int | None = None) -> list:
    """The master program's steps, or (with program_id) that client's copy."""
    owner = ProgramStep.program_id.is_(None) if program_id is None else ProgramStep.program_id == program_id
    return session.exec(select(ProgramStep).where(owner).order_by(ProgramStep.sort_order, ProgramStep.id)).all()


_COPIED_STEP_FIELDS = ("sort_order", "anchor", "offset_days", "section", *TEXT_FIELDS)
_COPIED_LINE_FIELDS = ("product_id", "animal_group", "dose", "note", "category", "fixed_quantity",
                       "quantity_override", "unit_price", "discount_percent")


def copy_master_into(session: Session, program: HerdingProgram, done: dict | None = None) -> None:
    """Gives the client their own copy of the master program. `done` maps a
    master step id to when that step was ticked off, to carry ticks over."""
    done = done or {}
    masters = ordered_steps(session)
    lines = {}
    for line in session.exec(select(ProgramStepProduct).where(
            ProgramStepProduct.step_id.in_([s.id for s in masters]))).all() if masters else []:
        lines.setdefault(line.step_id, []).append(line)
    for master in masters:
        step = ProgramStep(program_id=program.id, source_step_id=master.id, done_at=done.get(master.id),
                           **{f: getattr(master, f) for f in _COPIED_STEP_FIELDS})
        session.add(step)
        session.flush()
        for line in sorted(lines.get(master.id, []), key=lambda l: l.id):
            session.add(ProgramStepProduct(step_id=step.id, source_line_id=line.id,
                                           **{f: getattr(line, f) for f in _COPIED_LINE_FIELDS}))


def ensure_client_copy(session: Session, program: HerdingProgram) -> None:
    """A client's program gets its copy the first time it's used; ticks made
    on the old shared master (before copies existed) come along."""
    if program.mating_date is None or ordered_steps(session, program.id):
        return
    legacy = session.exec(select(ProgramStepProgress).where(ProgramStepProgress.program_id == program.id)).all()
    copy_master_into(session, program, {p.step_id: p.done_at for p in legacy})
    for row in legacy:
        session.delete(row)
    session.commit()


def own_step(session: Session, program: HerdingProgram, step_id: int):
    """The client's step by its id - or by the master step it was copied
    from (links made before every client had a copy)."""
    step = session.get(ProgramStep, step_id)
    if step is not None and step.program_id == program.id:
        return step
    return session.exec(select(ProgramStep).where(ProgramStep.program_id == program.id,
                                                  ProgramStep.source_step_id == step_id)).first()


def own_line(session: Session, program: HerdingProgram, line_id: int):
    line = session.get(ProgramStepProduct, line_id)
    if line is not None:
        step = session.get(ProgramStep, line.step_id)
        if step is not None and step.program_id == program.id:
            return line
    return session.exec(select(ProgramStepProduct).join(ProgramStep, ProgramStep.id == ProgramStepProduct.step_id)
                        .where(ProgramStep.program_id == program.id,
                               ProgramStepProduct.source_line_id == line_id)).first()


def step_products(session: Session, step_ids) -> dict:
    """{step_id: [(ProgramStepProduct, Product), ...]}"""
    out = {}
    if not step_ids:
        return out
    rows = session.exec(select(ProgramStepProduct, Product)
                        .where(ProgramStepProduct.step_id.in_(step_ids))
                        .where(ProgramStepProduct.product_id == Product.id)
                        .order_by(ProgramStepProduct.id)).all()
    for line, product in rows:
        out.setdefault(line.step_id, []).append((line, product))
    return out


def _status(day: date, done: bool, today: date) -> str:
    if done:
        return "done"
    if day < today:
        return "overdue"
    return "due" if day <= today + timedelta(days=DUE_WITHIN_DAYS) else "upcoming"


def step_date(step: ProgramStep, anchors: dict):
    if step.date_override is not None:
        return step.date_override.date() if isinstance(step.date_override, datetime) else step.date_override
    if step.anchor == UNDATED:
        return None
    return anchors.get(step.anchor, anchors["mating_start"]) + timedelta(days=step.offset_days)


def step_labels(session: Session, program: HerdingProgram, step_ids) -> dict:
    """{step id: (sort key, "Mon 25 May 2026 - stage")} for a program
    quote's sections, in the program's date order (undated last)."""
    if not step_ids:
        return {}
    anchors = anchor_dates(program, get_template(session)) if program.mating_date else None
    out = {}
    for step in session.exec(select(ProgramStep).where(ProgramStep.id.in_(step_ids))).all():
        day = step_date(step, anchors) if anchors else None
        when = f"{day:%a %d %b %Y}" if day else "Any time"
        out[step.id] = ((day is None, day or date.max, step.sort_order, step.id),
                        " - ".join(v for v in (when, step.stage) if v))
    return out


def schedule(session: Session, program: HerdingProgram, today: date | None = None) -> dict:
    """The client's own program: every step with its date and whether it's
    done/overdue/due/upcoming, each product line's amounts and costs for
    their animals (like the cost sheet), step subtotals and the totals."""
    empty = {"anchors": None, "steps": [], "sections": [], "progress": {"done": 0, "total": 0},
             "totals": {"cost_used": 0, "cost_buy": 0, "animals": 0, "per_animal": 0}}
    if program.mating_date is None:
        return empty
    ensure_client_copy(session, program)
    today = today or date.today()
    template = get_template(session)
    settings = session.get(BusinessSettings, 1) or BusinessSettings()
    anchors = anchor_dates(program, template)
    groups = session.exec(select(AnimalGroup).where(AnimalGroup.program_id == program.id)).all()
    steps = ordered_steps(session, program.id)
    products = step_products(session, [s.id for s in steps])
    out = []
    for step in steps:
        day = step_date(step, anchors)
        lines = []
        for line, product in products.get(step.id, []):
            head = headcount(groups, line.animal_group)
            # What's agreed with the client (shared with the program's quote) wins.
            price = line.unit_price if line.unit_price is not None else price_excl_vat(product, settings)
            costs = line_costs(product, line.dose, head, line.fixed_quantity, price)
            if line.quantity_override is not None:
                costs["buy"], costs["cost_buy"] = line.quantity_override, round(line.quantity_override * price, 2)
            if line.discount_percent:
                keep = 1 - line.discount_percent / 100
                costs["cost_used"], costs["cost_buy"] = round(costs["cost_used"] * keep, 2), round(costs["cost_buy"] * keep, 2)
            lines.append({**line.model_dump(), "product_name": product.name, "unit": product.unit,
                          "pack_size": product.pack_size, "price": product.price, "price_excl_vat": price,
                          "head": head, **costs})
        status = "none" if day is None else _status(day, step.done_at is not None, today)
        section = step.section or derive_section([line for line, _ in products.get(step.id, [])])
        if section is None and lines:
            section = OTHER_SECTION
        out.append({**step.model_dump(), "section": section, "date": day, "status": status, "products": lines,
                    "subtotal": round(sum(l["cost_used"] for l in lines), 2),
                    "subtotal_buy": round(sum(l["cost_buy"] for l in lines), 2)})
    # Dated steps in date order (a client's own dates can reorder them), then the medicine box.
    out.sort(key=lambda s: (s["date"] is None, s["date"] or date.max, s["sort_order"]))
    dated = [s for s in out if s["date"] is not None]
    animals = sum(g.group_size for g in groups)
    cost_used = round(sum(s["subtotal"] for s in out), 2)
    # The sheet's sections, each with its TOTAAL; a step with neither products
    # nor a section is a note from the Kudde program, shown by its date.
    sections = []
    for name in (*SECTIONS[:3], OTHER_SECTION, SECTIONS[3]):
        members = [s for s in out if s["section"] == name]
        if members:
            sections.append({"name": name, "step_ids": [s["id"] for s in members],
                             "subtotal": round(sum(s["subtotal"] for s in members), 2),
                             "subtotal_buy": round(sum(s["subtotal_buy"] for s in members), 2)})
    return {"anchors": anchors, "steps": out, "sections": sections,
            "progress": {"done": sum(1 for s in dated if s["status"] == "done"), "total": len(dated)},
            "totals": {"cost_used": cost_used, "cost_buy": round(sum(s["subtotal_buy"] for s in out), 2),
                       "animals": animals, "per_animal": round(cost_used / animals, 2) if animals else 0,
                       "vat_rate": settings.vat_rate if settings.vat_registered else 0}}


# --- Importing the business's Excel sheet as the template ---

class TemplateImportError(ValueError):
    pass


_COLUMNS = {  # template field -> words its header starts with (Afrikaans and English)
    "date": ("datum", "date"),
    "stage": ("produksie", "stadium", "stage"),
    "management": ("kudde", "bestuur", "management", "herd"),
    "vaccinations": ("inenting", "ent", "vaccin"),
    "dosing": ("dosering", "doseer", "dosing", "dose"),
    "vitamins": ("vitamien", "vitamin", "spoor", "mineral"),
    "feeding": ("voeding", "feed"),
}
_CELL = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_WEANING_WORDS = ("speen", "wean")


def _text(value) -> str | None:
    """Tidies a cell: the sheet starts new lines in a cell with runs of
    spaces ("1  Rams in      2  Ewes out") - those become line breaks."""
    if value is None:
        return None
    text = str(value).replace("\r", "")
    text = re.sub(r"[ \t]{2,}(?=\d{1,2}[.,]?\s)|[ \t]{4,}", "\n", text)
    text = "\n".join(re.sub(r"\s+", " ", line).strip() for line in text.split("\n"))
    return text.strip() or None


def _find_inputs(ws):
    """(mating date cell, season weeks cell, weaning months cell) from the
    labelled input rows above the table."""
    found = {}
    for row in ws.iter_rows(min_row=1, max_row=15):
        label = " ".join(str(c.value) for c in row if isinstance(c.value, str)).lower()
        values = [c for c in row if c.value is not None and not isinstance(c.value, str)]
        if not values:
            continue
        cell = values[-1]
        if "paardatum" in label or "mating date" in label:
            found.setdefault("mating", cell)
        elif "paarseisoen" in label or "mating season" in label:
            found.setdefault("weeks", cell)
        elif "gespeen" in label or "wean" in label:
            found.setdefault("weaning", cell)
    if not isinstance(found.get("mating") and found["mating"].value, (datetime, date)):
        raise TemplateImportError("Couldn't find the first mating date (a row labelled 'Eerste paardatum' with a date)")
    return found


def _find_header(ws):
    for row in ws.iter_rows(min_row=1, max_row=30):
        cols = {}
        for cell in row:
            if isinstance(cell.value, str):
                head = cell.value.strip().lower()
                for field, words in _COLUMNS.items():
                    if field not in cols and head.startswith(words):
                        cols[field] = cell.column_letter
        if "date" in cols and len(cols) >= 3:
            return row[0].row, cols
    raise TemplateImportError("Couldn't find the table's header row (Datum, Produksie Stadium, Kudde-Bestuur, ...)")


def _linear_value(formula, row, rows, refs, env, seen=()):
    """The formula's value as days, with M = first mating day and W = season weeks."""
    if row in seen:
        raise TemplateImportError(f"Row {row}'s date refers back to itself")

    def replace(match):
        coord = f"{match.group(1)}{match.group(2)}"
        if coord in refs:
            return refs[coord]
        if match.group(1) == rows["_date_col"] and int(match.group(2)) in rows:
            return f"R{match.group(2)}"
        raise TemplateImportError(f"Row {row}'s date uses cell {coord}, which isn't the mating date, season or another date")

    expr = _CELL.sub(replace, formula.lstrip("=").replace("$", ""))

    def ev(node):
        if isinstance(node, ast.Expression):
            return ev(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return node.value
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            return ev(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
        if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult)):
            a, b = ev(node.left), ev(node.right)
            return a + b if isinstance(node.op, ast.Add) else a - b if isinstance(node.op, ast.Sub) else a * b
        if isinstance(node, ast.Name):
            if node.id in env:
                return env[node.id]
            if node.id.startswith("R"):
                other = int(node.id[1:])
                return _linear_value(rows[other]["formula"], other, rows, refs, env, seen + (row,))
        raise TemplateImportError(f"Row {row}'s date formula ({formula}) is more than simple adding and subtracting")

    try:
        return ev(ast.parse(expr, mode="eval"))
    except SyntaxError:
        raise TemplateImportError(f"Couldn't read row {row}'s date formula ({formula})")


def _date_rules(rows, refs):
    """(anchor, offset_days) per row, plus the gestation and weaning days the
    sheet implies. A row is timed from lambing if it is (or refers to) the
    row that's mating + ~5 months, from weaning if it is (or refers to) the
    row whose text mentions weaning, else from mating; adding the season
    length moves it to the end of mating/lambing."""
    kinds, rules, constants = {}, {}, {}

    def kind(row):
        if row in kinds:
            return kinds[row]
        info = rows[row]
        referenced = [int(n) for n in re.findall(r"R(\d+)", _CELL.sub(
            lambda m: f"R{m.group(2)}" if m.group(1) == rows["_date_col"] else "", info["formula"].replace("$", "")))]
        base = info["values"]
        if any(w in (info["words"] or "") for w in _WEANING_WORDS) and not any(kind(r) == "weaning" for r in referenced):
            result = "weaning"
            constants.setdefault("weaning_days", base[0])
        elif referenced:
            result = kind(referenced[0])
        elif base[1] == 0 and 100 <= base[0] <= 200:
            result = "lambing"
            constants.setdefault("gestation_days", base[0])
        else:
            result = "mating"
        kinds[row] = result
        return result

    for row in (r for r in rows if isinstance(r, int)):
        kind(row)
    for row in (r for r in rows if isinstance(r, int)):
        days, season = rows[row]["values"]
        if season not in (0, 7):
            raise TemplateImportError(f"Row {row}'s date adds the season length in a way this app can't follow")
        k = kinds[row]
        if k == "mating":
            rules[row] = ("mating_end" if season else "mating_start", days)
        elif k == "lambing":
            rules[row] = ("lambing_end" if season else "lambing_start", days - constants["gestation_days"])
        else:
            if season:
                raise TemplateImportError(f"Row {row}'s date adds the mating season to the weaning day")
            rules[row] = ("weaning", days - constants["weaning_days"])
    return rules, constants


def parse_template_workbook(data: bytes) -> dict:
    try:
        wb = load_workbook(io.BytesIO(data))
    except Exception as exc:
        raise TemplateImportError("That file isn't an Excel workbook (.xlsx) this app can read") from exc
    ws = wb.worksheets[0]
    inputs = _find_inputs(ws)
    header_row, cols = _find_header(ws)
    refs = {inputs["mating"].coordinate: "M"}
    if "weeks" in inputs:
        refs[inputs["weeks"].coordinate] = "W"

    rows, name, current = {"_date_col": cols["date"]}, None, None
    for r in range(header_row + 1, ws.max_row + 1):
        first = ws[f"{cols['date']}{r}"].value
        texts = {f: ws[f"{c}{r}"].value for f, c in cols.items() if f != "date"}
        if isinstance(first, str) and first.startswith("="):
            current = r
            rows[r] = {"formula": first, "texts": {f: [] for f in TEXT_FIELDS}}
        elif isinstance(first, str) and first.strip() and not any(texts.values()):
            name = name or _text(first)  # a section heading such as "EERSTE LAMSEISOEN"
            continue
        if current is None:
            continue
        for field, value in texts.items():
            if _text(value):
                rows[current]["texts"][field].append(_text(value))
    steps = [r for r in rows if isinstance(r, int)]
    if not steps:
        raise TemplateImportError("Couldn't find any dates worked out from the mating date under the header row")
    for r in steps:
        rows[r]["words"] = " ".join(t for ts in rows[r]["texts"].values() for t in ts).lower()
        base = _linear_value(rows[r]["formula"], r, rows, refs, {"M": 0, "W": 0})
        with_season = _linear_value(rows[r]["formula"], r, rows, refs, {"M": 0, "W": 1})
        if _linear_value(rows[r]["formula"], r, rows, refs, {"M": 1000, "W": 0}) - base != 1000:
            raise TemplateImportError(f"Row {r}'s date isn't the mating date plus or minus some days")
        rows[r]["values"] = (int(round(base)), int(round(with_season - base)))
    rules, constants = _date_rules(rows, refs)

    def number(key, low, high):
        cell = inputs.get(key)
        value = cell.value if cell is not None else None
        return int(value) if isinstance(value, (int, float)) and low <= value <= high else None

    generic_title = re.fullmatch(r"(sheet|blad)\s*\d*", ws.title.strip(), re.I)
    return {
        "name": (name if generic_title else _text(ws.title)) or name,
        "settings": {k: v for k, v in {
            "mating_weeks": number("weeks", 1, 52), "weaning_months": number("weaning", 1, 12),
            "gestation_days": constants.get("gestation_days"), "weaning_days": constants.get("weaning_days"),
        }.items() if v is not None},
        "steps": [{"anchor": rules[r][0], "offset_days": rules[r][1],
                   **{f: "\n".join(rows[r]["texts"][f]) or None for f in TEXT_FIELDS}} for r in steps],
    }


def import_template(session: Session, data: bytes) -> dict:
    """Replaces the master program with the sheet's. A step with the same
    date rule as an existing one keeps its id - and so its products and
    every client's tick against it; steps no longer in the sheet go."""
    parsed = parse_template_workbook(data)
    template = get_template(session)
    template.name = (parsed["name"] or template.name)[:200]
    for key, value in parsed["settings"].items():
        setattr(template, key, value)
    if "weaning_days" in parsed["settings"]:
        template.weaning_rule = "fixed"  # the sheet's own rule; each program can switch to by-age
    template.updated_at = datetime.utcnow()
    session.add(template)

    existing = {}
    for step in ordered_steps(session):
        existing.setdefault((step.anchor, step.offset_days), []).append(step)
    created = updated = 0
    for order, spec in enumerate(parsed["steps"], start=1):
        matches = existing.get((spec["anchor"], spec["offset_days"]))
        step = matches.pop(0) if matches else None
        if step is None:
            step, created = ProgramStep(), created + 1
        else:
            updated += 1
        step.sort_order = order
        step.origin = "kudde"  # this sheet's now, even if the cost sheet made it
        for key, value in spec.items():
            setattr(step, key, value)
        session.add(step)
    # Steps only the cost sheet made (see import_cost_sheet) aren't this sheet's to remove.
    leftover = [s for ss in existing.values() for s in ss if s.origin != "cost"]
    for step in leftover:
        delete_step(session, step)
    session.commit()
    return {"created": created, "updated": updated, "removed": len(leftover), "name": template.name}


def delete_step(session: Session, step: ProgramStep) -> None:
    for model in (ProgramStepProduct, ProgramStepProgress):
        for row in session.exec(select(model).where(model.step_id == step.id)).all():
            session.delete(row)
    session.delete(step)


# --- Importing the business's cost sheet ("Ent en doseer kostes") ---
#
# One sheet: the headcount rows (TOTAAL OOIE / RAMME / LAMMERS / JONG
# OOITJIES / JONG RAMMETJIES), DEKTYD (the first mating day) and LAMTYD
# (= DEKTYD + gestation), then blocks of product lines, each block dated by
# a formula off DEKTYD or LAMTYD in column A. A line's "PRODUK TOTAAL"
# formula names the headcount cells it multiplies - which says which
# groups it's for - or just repeats the quantity (the medicine box).

def _tidy(text) -> str | None:
    """"(VOOR DEK ENTING)" -> "Voor dek enting"; mixed-case text is kept."""
    if not isinstance(text, str):
        return None
    text = re.sub(r"\s+", " ", text).strip().strip("!:").strip()
    if text.startswith("(") and text.endswith(")"):
        text = text[1:-1].strip()
    text = text.strip("!:").strip()
    if not text:
        return None
    return text[0].upper() + text[1:].lower() if text.upper() == text else text


def _group_label(text) -> str | None:
    t = (text or "").lower()
    if "jong" in t and "ooi" in t:
        return "Jong ooitjies"
    if "jong" in t and "ram" in t:
        return "Jong rammetjies"
    if "ooi" in t:
        return "Ooie"
    if "ram" in t:
        return "Ramme"
    if "lam" in t:
        return "Lammers"
    return None


def _section_heading(text) -> str | None:
    """"JONG OOITJIES EN RAMME", "DOSERING OOIE EN RAMME", "Medisyne Boks:"
    -> the section a heading starts."""
    t = str(text or "").lower()
    if "medisyne" in t:
        return SECTIONS[3]
    if "jong" in t:
        return SECTIONS[1]
    if "ooie" in t and "ram" in t:
        return SECTIONS[2]
    if t.strip(" :") == "lammers":
        return SECTIONS[0]
    return None


_FIXED_FORMULA = re.compile(r"=\s*(IFERROR\(\s*)?\$?F\$?\d+\s*[,)]?", re.IGNORECASE)


def _name_key(name) -> str:
    return re.sub(r"[^a-z0-9%]", "", str(name or "").lower())


def parse_cost_sheet(data: bytes, products) -> dict:
    """Blocks of product lines with their date rule, groups and doses.
    `products` are the app's products, matched by name (ignoring spaces
    and punctuation); unmatched names are listed, not guessed."""
    try:
        ws = load_workbook(io.BytesIO(data)).worksheets[0]
    except Exception as exc:
        raise TemplateImportError("That file isn't an Excel workbook (.xlsx) this app can read") from exc
    cell = lambda col, row: ws[f"{col}{row}"].value
    group_rows, mating_row, header_row = {}, None, None
    for r in range(1, min(ws.max_row, 40) + 1):
        a, b = cell("A", r), cell("B", r)
        if (header_row is None and isinstance(a, (int, float)) and not isinstance(a, bool)
                and isinstance(b, str) and _group_label(b)):
            group_rows[r] = _group_label(b)
        texts = " ".join(str(cell(c, r)).lower() for c in "ABCDEFGH" if isinstance(cell(c, r), str))
        if header_row is None and "produk" in texts and "doser" in texts:
            header_row = r
        if mating_row is None and isinstance(a, (datetime, date)) and isinstance(b, str) and "dek" in b.lower():
            mating_row = r
    if mating_row is None or header_row is None or not group_rows:
        raise TemplateImportError("That doesn't look like the cost sheet (it needs the headcount rows, the "
                                  "PRODUK / DOSERING header and DEKTYD with the first mating date)")

    formulas = {r: cell("A", r) for r in range(1, ws.max_row + 1)
                if ws[f"A{r}"].data_type == "f"}
    lambing_row = next((r for r, f in formulas.items() if re.fullmatch(rf"=\+?\$?A\$?{mating_row}\+(\d+)", f.replace(" ", ""))
                        and "lam" in str(cell("B", r) or "").lower()), None)
    gestation = int(re.search(r"\+(\d+)$", formulas[lambing_row].replace(" ", "")).group(1)) if lambing_row else 147

    def days(row, seen=()):
        """(days after the first mating day, whether it's counted from lambing)"""
        if row == mating_row:
            return 0, False
        if row in seen or row not in formulas:
            raise TemplateImportError(f"Couldn't work out the date in row {row}")
        refs = {}

        def swap(m):
            refs[f"R{m.group(1)}"] = int(m.group(1))
            return f"R{m.group(1)}"
        expr = re.sub(r"\$?A\$?(\d+)", swap, formulas[row].lstrip("=").replace("$", ""))
        values = {name: days(ref, seen + (row,)) for name, ref in refs.items()}
        via_lambing = row == lambing_row or any(v[1] for v in values.values())

        def ev(node):
            if isinstance(node, ast.Expression):
                return ev(node.body)
            if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
                return node.value
            if isinstance(node, ast.Name) and node.id in values:
                return values[node.id][0]
            if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
                return ev(node.operand) * (-1 if isinstance(node.op, ast.USub) else 1)
            if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult)):
                a, b = ev(node.left), ev(node.right)
                return a + b if isinstance(node.op, ast.Add) else a - b if isinstance(node.op, ast.Sub) else a * b
            raise TemplateImportError(f"Row {row}'s date formula is more than adding and subtracting days")
        try:
            return int(round(ev(ast.parse(expr, mode="eval")))), via_lambing
        except SyntaxError:
            raise TemplateImportError(f"Couldn't read row {row}'s date formula")

    by_name = {_name_key(p.name): p for p in products}
    blocks, block, unmatched = [], None, []
    section = SECTIONS[0]  # the sheet starts with the lambs, without a heading

    def open_block(anchor, offset, title):
        nonlocal block
        key = (section, anchor, offset)
        block = next((b for b in blocks if (b["section"], b["anchor"], b["offset_days"]) == key), None)
        if block is None:
            block = {"section": section, "anchor": anchor, "offset_days": offset, "stage": title,
                     "notes": [], "lines": []}
            blocks.append(block)
        elif title and not block["stage"]:
            block["stage"] = title
        elif title and title != block["stage"] and title not in block["notes"]:
            block["notes"].append(title)  # a second heading on the same day ("6-maande voor dek")

    for r in range(header_row + 1, ws.max_row + 1):
        if r in (mating_row, lambing_row):
            continue
        a, b, c, f, g = (cell(col, r) for col in "ABCFG")
        product_name = c.strip() if isinstance(c, str) and c.strip() else None
        label = str(b or "").strip().lower()
        if label.startswith("totaal") or label.startswith("totale"):
            continue
        if not product_name and _section_heading(b):
            section = _section_heading(b)
        if r in formulas:
            d, via_lambing = days(r)
            anchor, offset = ("lambing_start", d - gestation) if via_lambing else ("mating_start", d)
            open_block(anchor, offset, None if product_name else _tidy(b))
            if product_name and not block["stage"]:
                block["stage"] = _tidy(b)  # the date sits on the first product row
        elif "medisyne" in label and not product_name:
            open_block(UNDATED, 0, "Medisyne boks")
            continue
        if block is None:
            continue
        if isinstance(a, str) and not a.startswith("=") and not a.lower().startswith("verander"):
            note = _tidy(a)
            if note and note not in block["notes"]:
                block["notes"].append(note)
        if product_name is None:
            if isinstance(a, (datetime, date)) and _tidy(b) and _tidy(b) not in block["notes"]:
                block["notes"].append(_tidy(b))  # "Behandel jong ooitjies saam met ooie voor lam"
            continue
        product = by_name.get(_name_key(product_name))
        if product is None:
            if product_name not in unmatched:
                unmatched.append(product_name)
            continue
        refs = [int(n) for n in re.findall(r"\$?A\$?(\d+)", str(g or ""))]
        groups = list(dict.fromkeys(group_rows[n] for n in refs if n in group_rows))
        amount = f if isinstance(f, (int, float)) and not isinstance(f, bool) else None
        fixed = not groups and (block["anchor"] == UNDATED or bool(isinstance(g, str) and _FIXED_FORMULA.match(g.replace(" ", ""))))
        block["lines"].append({"product_id": product.id, "category": _tidy(b),
                               "animal_group": ", ".join(groups) or None,
                               "dose": None if fixed else amount, "fixed_quantity": amount if fixed else None})
    if not any(b["lines"] for b in blocks):
        raise TemplateImportError("Found no product lines with products this app knows - check the product names")
    for b in blocks:  # the animals on a block's lines say its section best; headings only fill gaps
        b["section"] = derive_section([SimpleNamespace(**l) for l in b["lines"]]) or b["section"]
    mating = cell("A", mating_row)
    return {"blocks": blocks, "unmatched": unmatched, "gestation": gestation,
            "mating_date": mating.date() if isinstance(mating, datetime) else mating,
            "counts": {label: int(cell("A", r)) for r, label in group_rows.items()}}


def _apply_blocks(session: Session, program_id: int | None, blocks, all_lines: bool) -> dict:
    """Puts parsed cost-sheet blocks into the master (program_id None) or one
    client's program. A block lands on the step with the same section and
    date rule - so a client's ticks and the step's id stay - else on a new
    "cost" step. Kudde program steps are never removed. `all_lines`: the
    sheet replaces every product line (a client's own sheet), not just the
    lines an earlier cost sheet made (the master)."""
    steps = ordered_steps(session, program_id)
    lines_by_step = {}
    for line in session.exec(select(ProgramStepProduct).where(
            ProgramStepProduct.step_id.in_([s.id for s in steps]))).all() if steps else []:
        lines_by_step.setdefault(line.step_id, []).append(line)
    by_key = {}
    for step in steps:
        section = step.section or derive_section(lines_by_step.get(step.id, []))
        if section and (step.origin == "cost" or all_lines):
            by_key.setdefault((section, step.anchor, step.offset_days), step)
    for step_lines in lines_by_step.values():
        for line in step_lines:
            if all_lines or line.origin == "cost":
                session.delete(line)
    session.flush()
    next_order = max((s.sort_order for s in steps), default=0) + 1
    created = lines = 0
    touched = set()
    for block in blocks:
        if not block["lines"] and not block["stage"]:
            continue
        key = (block["section"], block["anchor"], block["offset_days"])
        step = by_key.get(key)
        if step is None:
            step = ProgramStep(program_id=program_id, anchor=block["anchor"], offset_days=block["offset_days"],
                               stage=block["stage"], sort_order=next_order, origin="cost")
            next_order += 1
            created += 1
            session.add(step)
            session.flush()
            by_key[key] = step
        step.section = block["section"]
        if step.origin == "cost":
            step.stage = block["stage"] or step.stage
            step.management = "\n".join(block["notes"]) or None
        session.add(step)
        touched.add(step.id)
        for spec in block["lines"]:
            session.add(ProgramStepProduct(step_id=step.id, origin="cost", **spec))
            lines += 1
    session.flush()
    for step in ordered_steps(session, program_id):  # dates only an earlier sheet had
        if step.origin == "cost" and step.id not in touched:
            delete_step(session, step)
    return {"steps_created": created, "lines": lines}


def import_cost_sheet(session: Session, data: bytes) -> dict:
    """Puts the cost sheet's product lines into the master program, one step
    per block in its section, like the sheet (a Kudde program step on the
    same day stays a step of its own). Re-importing replaces the cost
    sheet's lines instead of adding them twice. Clients' own copies aren't
    touched - "Start again from the master" brings it in."""
    parsed = parse_cost_sheet(data, session.exec(select(Product)).all())
    result = _apply_blocks(session, None, parsed["blocks"], all_lines=False)
    session.commit()
    return {**result, "unmatched": parsed["unmatched"]}


def import_client_sheet(session: Session, program: HerdingProgram, data: bytes) -> dict:
    """A client's own cost sheet (one downloaded from the app, or the
    business's sheet filled in for them) into their program: DEKTYD, the
    animal numbers and every product line come from the sheet."""
    parsed = parse_cost_sheet(data, session.exec(select(Product)).all())
    mating = parsed["mating_date"]
    program.mating_date = datetime(mating.year, mating.month, mating.day)
    session.add(program)
    groups = {g.animal_type.strip().lower(): g for g in
              session.exec(select(AnimalGroup).where(AnimalGroup.program_id == program.id)).all()}
    for label, count in parsed["counts"].items():
        group = groups.get(label.lower())
        if group is None:
            group = AnimalGroup(program_id=program.id, animal_type=label, group_size=count)
        group.group_size = count
        session.add(group)
    session.flush()
    ensure_client_copy(session, program)
    result = _apply_blocks(session, program.id, parsed["blocks"], all_lines=True)
    session.commit()
    return {**result, "unmatched": parsed["unmatched"]}


# --- A client's program out as the business's cost sheet (.xlsx) ---

_COUNT_ROWS = {"ooie": 5, "ramme": 6, "lammers": 7, "jong ooitjies": 8, "jong rammetjies": 9}
_CLIENT_FIELDS = (("BESIGHEID", "farm_name"), ("NAAM", "name"), ("ADRES", "address"), ("E-POS", "email"),
                  ("SEL", "phone"), ("BTW NO", "vat_number"))


def export_client_sheet(session: Session, program: HerdingProgram, client, rep=(None, None, None)) -> bytes:
    """The program laid out like the cost sheet, with live formulas (counts
    in A5:A9, DEKTYD in A12, every date a formula off it, PRODUK TOTAAL and
    TOTAAL R worked out by Excel) - so it still works in Excel, and
    import_client_sheet reads it straight back."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    sched = schedule(session, program)
    if sched["anchors"] is None:
        raise TemplateImportError("Set the first mating date first")
    template = get_template(session)
    mating = sched["anchors"]["mating_start"]
    groups = {g.animal_type.strip().lower(): g.group_size for g in
              session.exec(select(AnimalGroup).where(AnimalGroup.program_id == program.id)).all()}
    bold, green = Font(bold=True), PatternFill("solid", fgColor="FF66FF33")
    money, day_fmt = '"R" #,##0.00', "dd mmmm yyyy"

    wb = Workbook()
    ws = wb.active
    ws.title = "ENT EN DOSEER KOSTE"
    for col, width in zip("ABCDEFGH", (23.6, 52.5, 44.6, 9.5, 14.5, 12.8, 12.8, 26)):
        ws.column_dimensions[col].width = width
    def text(ref, value):
        """Text someone typed into the app (names, notes, products): always a
        plain string cell - openpyxl would store "=..." as a live formula."""
        ws[ref] = value
        if isinstance(value, str):
            ws[ref].data_type = "s"

    for row, value in enumerate(rep, start=1):
        text(f"A{row}", value)
    ws["B1"] = f"VEEARTSENYPROGRAM KOSTE BEREKENING {sched['anchors']['lambing_start'].year}"
    ws["B1"].font = bold
    ws["A4"], ws["B4"] = "=SUM(A5:A9)", "TOTALE DIERE"
    for name, row in _COUNT_ROWS.items():
        label = next(g for g in GROUPS if g.lower() == name)
        ws[f"A{row}"], ws[f"B{row}"] = groups.get(name, 0), f"TOTAAL {label.upper()}"
        ws[f"A{row}"].font = bold
    for row, (label, field) in enumerate(_CLIENT_FIELDS, start=2):
        ws[f"C{row}"] = label
        text(f"D{row}", getattr(client, field, None) if client else None)
    for col, head in zip("ABCDEFGH", ("DATUM", "TYD", "PRODUK", "VERPAK", "PRYS EXCL VAT", "DOSERING ml / Dier",
                                      "PRODUK TOTAAL", "TOTAAL R EXCL VAT")):
        ws[f"{col}10"] = head
        ws[f"{col}10"].font, ws[f"{col}10"].fill = bold, green
    ws["A11"] = "VERANDER NET HIERDIE DATUM - ALTYD OP 'N MAANDAG"
    ws["A12"], ws["B12"] = datetime(mating.year, mating.month, mating.day), "DEKTYD"
    ws["A14"], ws["B14"] = f"=A12+{template.gestation_days}", "LAMTYD"
    for cell in ("A12", "A14"):
        ws[cell].number_format, ws[cell].font = day_fmt, bold

    def date_formula(step):
        if step["date"] is None:
            return None
        if step["date_override"] is None and step["anchor"] in ("mating_start", "lambing_start"):
            return f"={'A12' if step['anchor'] == 'mating_start' else 'A14'}{step['offset_days']:+d}"
        return f"=A12{(step['date'] - mating).days:+d}"  # ponytail: other anchors and typed-in dates come back as days from DEKTYD

    steps = {s["id"]: s for s in sched["steps"]}
    row, totals = 16, []
    for section in sched["sections"]:
        box = section["name"] == SECTIONS[3]
        ws[f"B{row}"] = "Medisyne Boks:" if box else section["name"].upper()
        ws[f"B{row}"].font, ws[f"B{row}"].fill = bold, green
        first = row + 1 if not box else row
        for step in (steps[i] for i in section["step_ids"]):
            if not box:
                row += 1
                ws[f"A{row}"] = date_formula(step)
                text(f"B{row}", step["stage"])
                ws[f"A{row}"].number_format, ws[f"A{row}"].font, ws[f"B{row}"].font = day_fmt, bold, bold
            notes = [n for n in (step["management"] or "").split("\n") if n.strip()] if step["origin"] == "cost" else []
            for i, line in enumerate(step["products"]):
                row += 1
                if i < len(notes):
                    text(f"A{row}", notes[i])
                text(f"B{row}", line["category"])
                text(f"C{row}", line["product_name"])
                ws[f"D{row}"], ws[f"E{row}"] = line["pack_size"], line["price_excl_vat"]
                if line["fixed_quantity"]:
                    ws[f"F{row}"], ws[f"G{row}"] = line["fixed_quantity"], f'=IFERROR(F{row},"")'
                else:
                    refs = [_COUNT_ROWS[n] for n in group_names(line["animal_group"]) if n in _COUNT_ROWS] \
                        or list(_COUNT_ROWS.values())
                    per_pack = f"/D{row}" if line["pack_size"] else ""
                    ws[f"F{row}"] = line["dose"]
                    ws[f"G{row}"] = f'=IFERROR(({"+".join(f"$A${r}" for r in refs)})*F{row}{per_pack},"")'
                ws[f"H{row}"] = f'=IFERROR(G{row}*E{row},"")'
                ws[f"E{row}"].number_format, ws[f"H{row}"].number_format = money, money
                ws[f"G{row}"].number_format = "0.00"
        row += 1
        ws[f"B{row}"], ws[f"H{row}"] = "TOTAAL", f"=SUM(H{first}:H{row - 1})"
        ws[f"B{row}"].font, ws[f"H{row}"].font, ws[f"H{row}"].number_format = bold, bold, money
        totals.append(f"H{row}")
        row += 2
    ws[f"B{row}"], ws[f"H{row}"] = "TOTALE KOSTE:", "=" + ("+".join(totals) or "0")
    ws[f"B{row}"].font, ws[f"H{row}"].font, ws[f"H{row}"].number_format = bold, bold, money
    ws[f"D{row + 2}"], ws[f"H{row + 2}"] = "KOSTE PER DIER PER JAAR", f"=IFERROR(H{row}/A4,0)"
    ws[f"D{row + 2}"].font, ws[f"H{row + 2}"].number_format = bold, money
    ws.print_area = f"A1:H{row + 2}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
