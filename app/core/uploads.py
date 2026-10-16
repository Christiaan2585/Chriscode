import csv
import io

from fastapi import HTTPException, UploadFile

IMPORT_LIMIT_BYTES = 10 * 1024 * 1024  # a price list or client list is a few hundred KB; this leaves room
# An .xlsx is a zip: 10 MB can unpack into millions of cells. Imports never need more than this.
MAX_TABLE_ROWS = 20_000
MAX_TABLE_COLUMNS = 100


async def read_capped(file: UploadFile, limit: int = IMPORT_LIMIT_BYTES) -> bytes:
    """The upload's bytes, or a 413 if it is bigger than `limit` (read one byte past it, never the whole thing)."""
    data = await file.read(limit + 1)
    if len(data) > limit:
        raise HTTPException(status_code=413, detail=f"That file is over {limit // (1024 * 1024)} MB - please use a smaller one")
    return data


def _text(value):
    """A cell as text, the way a person typed it: 12 not 12.0, blank = None."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def read_table(filename: str, data: bytes) -> tuple:
    """A CSV or .xlsx upload as (column names, rows as {column: text or None}). The first row is the header;
    unnamed columns and empty rows at the end are dropped. Everything stays text, so 0821234567 keeps its 0."""
    name = filename.lower()
    try:
        if name.endswith(".csv"):
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1252")  # Excel's "CSV" on Windows
            rows = csv.reader(io.StringIO(text))
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook  # only when a sheet is imported (keeps start-up quick)
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            rows = ws.iter_rows(values_only=True, max_col=MAX_TABLE_COLUMNS)
        else:
            raise HTTPException(status_code=400, detail="Old .xls files can't be read - open it in Excel and save it as .xlsx (or .csv)")
        grid = []
        for row in rows:  # read lazily, so an oversized sheet stops here instead of filling the memory
            if len(grid) > MAX_TABLE_ROWS:
                raise HTTPException(status_code=413, detail=f"That sheet has more than {MAX_TABLE_ROWS:,} rows - split it into smaller files")
            grid.append(list(row)[:MAX_TABLE_COLUMNS])
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the file - is it a real .xlsx or .csv file?")
    if not grid:
        return [], []
    header = [_text(h) for h in grid[0]]
    keep = [(i, h) for i, h in enumerate(header) if h]
    rows = [{h: _text(row[i]) if i < len(row) else None for i, h in keep} for row in grid[1:]]
    while rows and not any(rows[-1].values()):
        rows.pop()
    return [h for _, h in keep], rows
